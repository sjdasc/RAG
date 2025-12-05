import os
import glob
from pymilvus import connections, FieldSchema, CollectionSchema, DataType, Collection, utility
from txtai.embeddings import Embeddings
import numpy as np

# =============================================================================
# 설정
# =============================================================================
DATA_DIR = "./data"
MILVUS_HOST = "127.0.0.1"
MILVUS_PORT = "19530"
COLLECTION_NAME = "study_docs"
ALLOWED_EXT=["pdf","doc","docx","ppt","pptx","txt","md","html"]

# txtai 설정 (벡터 생성 엔진)
EMBEDDINGS_CONFIG = {
    "path": "sentence-transformers/all-MiniLM-L6-v2",
    "content": True,
}
EMBED_DIM = 384

# =============================================================================
# 1. txtai Embeddings 초기화
# =============================================================================
def initialize_embeddings():
    """txtai Embeddings 객체 생성 (벡터 생성 엔진)"""
    print("[TXTAI] Initializing txtai embeddings...")
    print(f"[TXTAI] Config: {EMBEDDINGS_CONFIG}")
    
    embeddings = Embeddings(EMBEDDINGS_CONFIG)
    
    print("[TXTAI] Model loaded successfully")
    print(f"[TXTAI] Embedding dimension: {EMBED_DIM}")
    return embeddings

# =============================================================================
# 2. Milvus 연결 및 컬렉션 설정
# =============================================================================
def connect_milvus():
    """Milvus Standalone에 연결"""
    print(f"\n[MILVUS] Connecting to {MILVUS_HOST}:{MILVUS_PORT}...")
    connections.connect("default", host=MILVUS_HOST, port=MILVUS_PORT)
    print("[MILVUS] Connected successfully")

def setup_milvus_collection(embed_dim: int = 384):
    """Milvus 컬렉션 생성"""
    
    # 기존 컬렉션 확인
    if COLLECTION_NAME in utility.list_collections():
        print(f"[MILVUS] Collection '{COLLECTION_NAME}' already exists.")
        collection = Collection(COLLECTION_NAME)
        collection.load()
        return collection
    
    # 새로 생성
    print(f"[MILVUS] Collection '{COLLECTION_NAME}' not found → creating new one.")
    
    # 스키마 정의
    fields = [
        FieldSchema(name="id", dtype=DataType.INT64, is_primary=True, auto_id=True),
        FieldSchema(name="vector", dtype=DataType.FLOAT_VECTOR, dim=embed_dim),
        FieldSchema(name="text", dtype=DataType.VARCHAR, max_length=65535),
        FieldSchema(name="path", dtype=DataType.VARCHAR, max_length=512),
        FieldSchema(name="filename", dtype=DataType.VARCHAR, max_length=256),
        FieldSchema(name="doc_type", dtype=DataType.VARCHAR, max_length=50),
    ]
    
    schema = CollectionSchema(fields, description="txtai → milvus integration")
    collection = Collection(COLLECTION_NAME, schema)
    
    # 인덱스 생성
    index_params = {
        "index_type": "IVF_FLAT",
        "metric_type": "IP",
        "params": {"nlist": 1024}
    }
    collection.create_index("vector", index_params)
    collection.load()
    
    print(f"[MILVUS] Collection '{COLLECTION_NAME}' created with schema:")
    print(f"  - Embedding dimension: {embed_dim}")
    print(f"  - Index type: IVF_FLAT")
    
    return collection

# =============================================================================
# 3. 파일에서 텍스트 추출
# =============================================================================
def extract_text(path: str) -> str:
    """다양한 형식의 파일에서 텍스트 추출"""
    lower = path.lower()
    try:
        if lower.endswith(".pdf"):
            from pypdf import PdfReader
            reader = PdfReader(path)
            return "\n".join(page.extract_text() or "" for page in reader.pages)
        
        elif lower.endswith((".docx", ".doc")):
            from docx import Document
            doc = Document(path)
            return "\n".join(p.text for p in doc.paragraphs)
        
        elif lower.endswith((".pptx", ".ppt")):
            from pptx import Presentation
            prs = Presentation(path)
            texts = []
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text:
                        texts.append(shape.text)
            return "\n".join(texts)
        
        else:  # txt 및 기타
            with open(path, "r", encoding="utf-8", errors="ignore") as f:
                return f.read()
    
    except Exception as e:
        print(f"[WARN] Failed to extract {path}: {e}")
        return ""

def load_all_documents(root_dir: str):
    print(f"\n[LOAD] Scanning {root_dir}...")

    # 1) Milvus 컬렉션 Load 보장 (중복 방지 확실하게)
    col = Collection(COLLECTION_NAME)
    col.load()

    try:
        existing = col.query(
            expr="id >= 0",
            output_fields=["filename", "path"]
        )
        existing_files = {(x["filename"], x["path"]) for x in existing}

    except:
        existing_files = set()

    documents=[]
    patterns=ALLOWED_EXT

    for pattern in patterns:
        for file_path in glob.glob(os.path.join(root_dir, pattern), recursive=True):
            filename = os.path.basename(file_path)

            # 🔥 filename + path 모두 검증하여 중복 방지
            if (filename, file_path) in existing_files:
                print(f"  [SKIP] Already indexed → {filename}")
                continue

            text = extract_text(file_path)
            if not text.strip(): 
                continue

            documents.append({
                "path": file_path,
                "filename": filename,
                "text": text,
                "doc_type": os.path.splitext(filename)[1][1:]
            })

            print(f"  [LOAD] NEW FILE → {filename} ({len(text)} chars)")

    print(f"[INFO] New documents loaded: {len(documents)}")
    return documents


# =============================================================================
# 4. ✅ numpy array 처리 수정 - txtai 벡터화 및 Milvus 저장
# =============================================================================
def vectorize_and_index_via_txtai(embeddings: Embeddings, collection: Collection, documents: list):
    """
    ✅ numpy array 처리 수정
    
    플로우:
    1. 문서 텍스트 추출
    2. txtai.batchtransform()으로 벡터화
    3. numpy array → 리스트 변환 (안전하게!)
    4. 벡터 → Milvus에 저장
    """
    if not documents:
        print("[ERROR] No documents to index")
        return
    
    print(f"\n[VECTORIZE] Processing {len(documents)} documents via txtai...")
    
    # ✅ Step 1: 문서 텍스트 추출
    texts = [doc["text"] for doc in documents]
    print(f"[VECTORIZE] Extracted {len(texts)} texts")
    
    # ✅ Step 2: txtai.batchtransform()으로 벡터화
    print("[VECTORIZE] Vectorizing via txtai.batchtransform()...")
    
    try:
        vectors = embeddings.batchtransform(texts)
        
        print(f"[VECTORIZE] Vectorization complete via txtai")
        
        # ✅ Step 3: numpy array 타입 확인 및 변환
        print(f"[VECTORIZE] Vectors type: {type(vectors)}")
        print(f"[VECTORIZE] Vectors shape: {vectors.shape if isinstance(vectors, np.ndarray) else 'N/A'}")
        
        # numpy array 확인
        if isinstance(vectors, np.ndarray):
            print(f"[VECTORIZE] Converting numpy array to list...")
            print(f"[VECTORIZE] Array shape: {vectors.shape} (rows, cols)")
            
            # ✅ 안전한 변환: numpy array를 명시적으로 리스트로
            vectors_list = vectors.tolist()
            print(f"[VECTORIZE] Converted to list with {len(vectors_list)} vectors")
            
            # 첫 벡터 확인
            if vectors_list and len(vectors_list) > 0:
                first_vec = vectors_list[0]
                if isinstance(first_vec, list):
                    print(f"[VECTORIZE] First vector dimension: {len(first_vec)}")
                    print(f"[VECTORIZE] First vector sample: {first_vec[:5]}...")
            
            vectors = vectors_list
        else:
            print(f"[VECTORIZE] Vectors is already list-like: {type(vectors)}")
        
        # ✅ 벡터 개수 확인 (안전하게 - numpy array bool 문제 해결)
        num_vectors = len(vectors)
        print(f"[VECTORIZE] Total vectors: {num_vectors}")
        
        # 벡터가 없으면 에러
        if num_vectors == 0:
            print("[ERROR] No vectors generated!")
            return
    
    except Exception as e:
        print(f"[ERROR] Vector extraction from txtai failed: {e}")
        import traceback
        traceback.print_exc()
        return
    
    # ✅ Step 4: Milvus에 삽입
    print(f"\n[MILVUS] Preparing {len(documents)} entities for Milvus insertion...")
    
    entities = []
    for i, (doc, vector) in enumerate(zip(documents, vectors)):
        # ✅ 벡터 타입 재확인 및 변환
        if isinstance(vector, np.ndarray):
            vector = vector.tolist()
        elif not isinstance(vector, list):
            vector = list(vector)
        
        # ✅ 벡터 길이 확인
        if len(vector) != EMBED_DIM:
            print(f"[WARN] Vector {i} has wrong dimension: {len(vector)} (expected {EMBED_DIM})")
            continue
        
        entity = {
            "vector": vector,
            "text": doc["text"][:65000],  # Milvus VARCHAR 제한
            "path": doc["path"],
            "filename": doc["filename"],
            "doc_type": doc["doc_type"]
        }
        entities.append(entity)
    
    if not entities:
        print("[ERROR] No valid entities to insert!")
        return
    
    # ✅ Milvus에 일괄 삽입
    print(f"[MILVUS] Inserting {len(entities)} entities into Milvus...")
    try:
        insert_result = collection.insert(entities)
        collection.flush()
        collection.load()
        
        print(f"[MILVUS] Successfully inserted {len(insert_result.primary_keys)} documents")
        print(f"[MILVUS] Collection now has {collection.num_entities} total entities")
    
    except Exception as e:
        print(f"[ERROR] Milvus insertion failed: {e}")
        import traceback
        traceback.print_exc()

# =============================================================================
# 5. 통계 및 검증
# =============================================================================
def show_collection_stats(collection):
    """컬렉션 통계 표시"""
    print(f"\n[STATS]")
    print(f"  - Collection name: {COLLECTION_NAME}")
    print(f"  - Total entities: {collection.num_entities}")
    print(f"  - Schema fields: {[field.name for field in collection.schema.fields]}")



# ============================================================
# 🔥 Milvus Collection 삭제 함수 - 주의!!!!
# ============================================================
def drop_milvus_collection_and_count():
    try:
        if COLLECTION_NAME in utility.list_collections():

            col = Collection(COLLECTION_NAME)
            count = col.num_entities  # 삭제 전 문서 수 확인

            utility.drop_collection(COLLECTION_NAME)
            print(f"🗑 Collection '{COLLECTION_NAME}' removed → {count} docs deleted")

            setup_milvus_collection(embed_dim=EMBED_DIM)
            print("📌 New empty collection initialized.")

            return count  # ← 삭제된 문서 수 반환
        
        else:
            print("⚠ No collection found. Nothing deleted.")
            return 0

    except Exception as e:
        print(f"❌ Error during clear: {e}")
        return 0

# ============================================================
# 🔥 파일명 기반 삭제 (문서개수/삭제된 ID 반환)
# ============================================================
def delete_document_by_filename(filename: str):
    try:
        col = Collection(COLLECTION_NAME)
        col.load()

        # filename 일치하는 데이터 조회
        matches = col.query(expr=f'filename == "{filename}"', output_fields=["id"])
        if not matches:
            print(f"⚠ No document found: {filename}")
            return 0, []

        ids = [m["id"] for m in matches]

        # 실제 DB 삭제
        col.delete(expr=f'id in {ids}')
        col.flush()
        print(f"🗑 Deleted {len(ids)} vectors from '{filename}'")

        # 로컬 파일도 제거
        local_path = f"./data/{filename}"
        if os.path.exists(local_path):
            os.remove(local_path)
            print(f"🗂 Local file removed: {local_path}")

        return len(ids), ids  # 삭제된 문서수 + ID 목록 반환

    except Exception as e:
        print(f"❌ Delete operation failed: {e}")
        return 0, []

